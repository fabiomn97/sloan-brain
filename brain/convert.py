#!/usr/bin/env python3
"""
Convert everything in raw/ and inbox/ into Markdown in corpus/.

Every output file starts with front matter (course, term, type, module, date,
source) and keeps page and slide numbers as "## Page 7" / "## Slide 12 — Title"
headings, so answers can cite exactly where something came from.

Inbox layout (drop anything here, any time):

  inbox/canvas/<course>/       files the harvester couldn't get, e.g. inbox/canvas/15.010/
  inbox/transcripts/           coffee chats, meetings (.txt .md .vtt .srt .docx .pdf)
  inbox/newsletters/           Sloan newsletters (.pdf .html .eml .txt)
  inbox/notes/                 your own notes, anything else

Start a file name with a date to date it: "2026-10-03 Coffee chat with Ana.txt".

Needs: pip install -r requirements.txt  (pymupdf, python-pptx, python-docx, openpyxl)

    python brain/convert.py            convert new and changed files
    python brain/convert.py --all      reconvert everything
"""

import argparse
import csv
import email
import email.policy
import email.utils
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from html.parser import HTMLParser

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (CORPUS_DIR, INBOX_DIR, RAW_DIR, ROOT, file_sha1,  # noqa: E402
                    load_manifest, read_doc, slugify, write_doc)

STATE = os.path.join(RAW_DIR, "convert-state.json")
CONVERT_FAILED = os.path.join(RAW_DIR, "convert-failed.csv")
DATED = re.compile(r"^(\d{4}-\d{2}-\d{2})[\s_-]+(.*)$")

KIND_TO_TYPE = {"syllabus": "syllabus", "page": "page", "assignment": "assignment",
                "announcement": "announcement", "discussion": "discussion"}
INBOX_TYPES = {"transcripts": "transcript", "newsletters": "newsletter", "notes": "note"}
PLURAL = {"slides": "slides", "reading": "readings", "case": "cases", "syllabus": "syllabus",
          "assignment": "assignments", "page": "pages", "announcement": "announcements",
          "discussion": "discussions", "transcript": "transcripts", "newsletter": "newsletters",
          "note": "notes", "document": "documents", "spreadsheet": "spreadsheets"}


BULLET_CHARS = "•▪●○■□◦‣∙➢➤✓\uf0b7\uf0a7\uf0d8\uf06e\uf0fc"
INLINE_BULLET = re.compile(r"\s[" + BULLET_CHARS + r"]\s")


def _block_text(raw):
    """One PDF text block -> a paragraph, or list items when it holds bullets."""
    raw = re.sub(r"-\n(?=[a-z])", "", raw.replace("\u00ad", ""))   # re-join hyphenated words
    raw = INLINE_BULLET.sub("\n\u2022 ", raw)
    out = []
    for line in raw.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line[0] in BULLET_CHARS or line.startswith(("- ", "\u2013 ")):
            out.append("- " + line.lstrip(BULLET_CHARS + "-\u2013 ").strip())
        elif out:
            out[-1] += " " + line
        else:
            out.append(line)
    return "\n".join(out).strip()


class Unsupported(Exception):
    pass


# --------------------------------------------------------------------------
# HTML to Markdown (Canvas pages, newsletters)
# --------------------------------------------------------------------------


class _HtmlToMd(HTMLParser):
    BLOCK = {"p", "div", "section", "article", "header", "footer", "tr", "blockquote",
             "h1", "h2", "h3", "h4", "h5", "h6", "li", "ul", "ol", "table", "pre", "br", "hr"}
    SKIP = {"script", "style", "iframe", "noscript", "svg", "head", "button", "form"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip, self.lists, self.href, self.link_text = [], 0, [], None, []
        self.meta, self.cell, self.marks = {}, False, []

    def _nl(self, n=2):
        tail = "".join(self.out[-3:])
        need = n - (len(tail) - len(tail.rstrip("\n")))
        if self.out and need > 0:
            self.out.append("\n" * need)

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "meta" and a.get("name"):
            self.meta[a["name"]] = a.get("content", "")
        if tag in self.SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._nl()
            # Keep "##" reserved for page/slide markers: page headings start at ###.
            self.out.append("#" * min(6, int(tag[1]) + 2) + " ")
        elif tag in ("p", "div", "section", "blockquote", "table", "pre"):
            self._nl()
        elif tag in ("ul", "ol"):
            self._nl(1)
            self.lists.append([tag, 0])
        elif tag == "li":
            self._nl(1)
            depth = max(0, len(self.lists) - 1)
            if self.lists and self.lists[-1][0] == "ol":
                self.lists[-1][1] += 1
                self.out.append("  " * depth + f"{self.lists[-1][1]}. ")
            else:
                self.out.append("  " * depth + "- ")
        elif tag == "br":
            self.out.append("\n")
        elif tag == "tr":
            self._nl(1)
            self.out.append("|")
        elif tag in ("td", "th"):
            self.cell = True
            self.out.append(" ")
        elif tag in ("strong", "b", "em", "i"):
            self.marks.append((tag, len(self.out)))
        elif tag == "a":
            self.href, self.link_text = a.get("href"), []

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in ("ul", "ol") and self.lists:
            self.lists.pop()
            self._nl()
        elif tag in ("td", "th"):
            self.out.append(" |")
            self.cell = False
        elif tag in ("strong", "b", "em", "i") and self.marks:
            _, start = self.marks.pop()
            inner = "".join(self.out[start:])
            del self.out[start:]
            mark = "**" if tag in ("strong", "b") else "*"
            core = inner.strip()
            if core and "\n" not in core:
                lead, trail = inner[:len(inner) - len(inner.lstrip())], inner[len(inner.rstrip()):]
                self.out.append(f"{lead}{mark}{core}{mark}{trail}")
            else:
                self.out.append(inner)
        elif tag == "a" and self.href is not None:
            text = "".join(self.link_text).strip()
            href = self.href
            self.href = None
            if href.startswith("http") and text and text != href and "/files/" not in href:
                self.out.append(f" ({href})")
        elif tag == "li":
            self._nl(1)
        elif tag in self.BLOCK:
            self._nl()

    def handle_data(self, data):
        if self.skip:
            return
        text = re.sub(r"\s+", " ", data)
        if self.href is not None:
            self.link_text.append(text)
        if text.strip() or (self.out and not self.out[-1].endswith((" ", "\n"))):
            self.out.append(text)

    def markdown(self):
        text = "".join(self.out)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


def html_to_md(raw):
    p = _HtmlToMd()
    p.feed(raw)
    p.close()
    return p.markdown(), p.meta


# --------------------------------------------------------------------------
# File converters. Each returns (body_markdown, info dict).
# --------------------------------------------------------------------------


def _clean_pages(pages):
    """Drop running headers/footers that repeat on most pages, and bare page numbers."""
    if len(pages) >= 3:
        counts = {}
        for text in pages:
            for line in {ln.strip() for ln in text.splitlines() if 0 < len(ln.strip()) <= 90}:
                counts[line] = counts.get(line, 0) + 1
        repeated = {ln for ln, n in counts.items() if n >= max(3, len(pages) * 0.5)}
    else:
        repeated = set()
    out = []
    for text in pages:
        lines = [ln for ln in text.splitlines()
                 if ln.strip() not in repeated and not re.fullmatch(r"\s*\d{1,3}\s*", ln)]
        out.append("\n".join(lines).strip())
    return out


def convert_pdf(path):
    import pymupdf
    doc = pymupdf.open(path)
    pages, landscape = [], 0
    for page in doc:
        if page.rect.width > page.rect.height:
            landscape += 1
        paras = []
        for block in page.get_text("blocks", sort=True):
            if block[6] != 0:          # image block
                continue
            text = _block_text(block[4])
            if text:
                paras.append(text)
        pages.append("\n\n".join(paras))
    n = max(1, len(pages))
    pages = _clean_pages(pages)
    chars = sum(len(p) for p in pages)
    if chars / n < 25:
        raise Unsupported("scanned PDF with no text layer (needs OCR)")
    words_per_page = sum(len(p.split()) for p in pages) / n
    body = "\n\n".join(f"## Page {i}\n\n{t}" for i, t in enumerate(pages, 1) if t)
    looks_like_slides = landscape / n > 0.7 and words_per_page < 160
    return body, {"locator_kind": "page", "slides": looks_like_slides,
                  "title": (doc.metadata or {}).get("title", "")}


def _shape_text(shape, out):
    if shape.shape_type == 6 and hasattr(shape, "shapes"):          # group
        for s in shape.shapes:
            _shape_text(s, out)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for r in shape.table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in r.cells]
            out.append("| " + " | ".join(cells) + " |")
        out.append("")
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip() or para.text.strip()
            if text:
                out.append("  " * para.level + "- " + text)


def convert_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip().replace("\n", " ") if title_shape is not None else ""
        lines = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            _shape_text(shape, lines)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        part = f"## Slide {i}" + (f" — {title}" if title else "") + "\n\n" + "\n".join(lines)
        if notes:
            part += "\n\n**Speaker notes:** " + re.sub(r"\s*\n\s*", " ", notes)
        slides.append(part.strip())
    if not any(len(s) > 20 for s in slides):
        raise Unsupported("deck has no extractable text (images only)")
    return "\n\n".join(slides), {"locator_kind": "slide", "slides": True}


def convert_docx(path):
    import docx
    d = docx.Document(path)
    out = []
    body = d.element.body
    paras = {p._element: p for p in d.paragraphs}
    tables = {t._element: t for t in d.tables}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading"):
                level = re.findall(r"\d", style)
                out.append("#" * min(6, (int(level[0]) if level else 1) + 2) + " " + text)
            elif style == "title":
                out.append("### " + text)
            elif "list" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif el in tables:
            for r in tables[el].rows:
                out.append("| " + " | ".join(c.text.replace("\n", " ").strip() for c in r.cells) + " |")
    text = "\n\n".join(out)
    if len(text) < 20:
        raise Unsupported("document is empty")
    return text, {"locator_kind": "section"}


def convert_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        # Hidden sheets are tool scratch space (e.g. Analytic Solver's rsklib* data), not content.
        if ws.sheet_state != "visible" or ws.title.lower().startswith("rsklib"):
            continue
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [clean_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append("| " + " | ".join(cells[:20]) + " |")
                if len(rows) == 1:
                    rows.append("|" + " --- |" * min(20, len(cells)))
            if len(rows) >= 150:
                rows.append("*(sheet truncated after 150 rows)*")
                break
        if rows:
            parts.append(f"## Sheet {ws.title}\n\n" + "\n".join(rows))
    if not parts:
        raise Unsupported("spreadsheet is empty")
    return "\n\n".join(parts), {"locator_kind": "sheet"}


JUNK = re.compile(r"_x[0-9A-Fa-f]{4}_|[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean_cell(value, limit=300):
    """Spreadsheet cell -> short plain text. Drops encoded binary that some add-ins store in cells."""
    if value is None:
        return ""
    text = JUNK.sub("", str(value)).replace("\n", " ").replace("|", "/").strip()
    if text and sum(not (c.isprintable() or c.isspace()) for c in text) > len(text) * 0.1:
        return ""
    return text if len(text) <= limit else text[:limit] + "…"


def convert_csv(path):
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        rows = list(csv.reader(f))[:150]
    return "\n".join("| " + " | ".join(r[:20]) + " |" for r in rows), {"locator_kind": "section"}


def convert_html(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        body, meta = html_to_md(f.read())
    # Our own saved pages repeat the title as <h1>; the front matter already has it.
    body = re.sub(r"^### [^\n]*\n+", "", body, count=1) if meta.get("kind") else body
    return body, {"locator_kind": "section", "meta": meta}


def convert_text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    meta, body = read_doc(path) if text.startswith("---\n") else ({}, text)
    if path.lower().endswith((".vtt", ".srt")):
        lines = []
        for line in body.splitlines():
            line = line.strip()
            if (not line or line == "WEBVTT" or line.isdigit() or "-->" in line
                    or line.startswith(("NOTE", "Kind:", "Language:"))):
                continue
            if not lines or lines[-1] != line:
                lines.append(line)
        body = "\n".join(lines)
    return body, {"locator_kind": "section", "meta": meta}


def convert_eml(path):
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=email.policy.default)
    part = msg.get_body(preferencelist=("html", "plain"))
    content = part.get_content() if part else ""
    body = html_to_md(content)[0] if part and part.get_content_type() == "text/html" else content
    return body, {"locator_kind": "section", "title": msg.get("subject", ""),
                  "date_header": msg.get("date", "")}


def convert_legacy_office(path):
    """.doc / .ppt / .xls: convert with LibreOffice when it's installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise Unsupported("old Office format; install LibreOffice or re-save as .docx/.pptx/.xlsx")
    target = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}[os.path.splitext(path)[1].lower()]
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", target, "--outdir", tmp, path],
                       check=True, capture_output=True, timeout=180)
        out = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + "." + target)
        return CONVERTERS["." + target](out)


CONVERTERS = {
    ".pdf": convert_pdf, ".pptx": convert_pptx, ".docx": convert_docx, ".xlsx": convert_xlsx,
    ".xlsm": convert_xlsx, ".csv": convert_csv, ".html": convert_html, ".htm": convert_html,
    ".txt": convert_text, ".md": convert_text, ".vtt": convert_text, ".srt": convert_text,
    ".eml": convert_eml, ".doc": convert_legacy_office, ".ppt": convert_legacy_office,
    ".xls": convert_legacy_office,
}


# --------------------------------------------------------------------------
# Classification
# --------------------------------------------------------------------------


def infer_type(name, folder="", module="", ext="", info=None):
    hay = f"{name} {folder} {module}".lower()
    if "syllab" in hay:
        return "syllabus"
    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        return "spreadsheet"
    if re.search(r"\bcases?\b|\bhbs\b|harvard business", hay):
        return "case"
    if re.search(r"\b(problem set|pset|homework|hw\s?\d|assignment|deliverable|template)\b", hay):
        return "assignment"
    if ext in (".pptx", ".ppt") or (info or {}).get("slides"):
        return "slides"
    if ext == ".pdf" and re.search(r"\b(slides?|lecture|deck)\b", hay):
        return "slides"
    if re.search(r"\b(readings?|article|chapter|ch\.?\s?\d|paper|note)\b", hay) or ext == ".pdf":
        return "reading"
    return "document"


def clean_title(name):
    base = os.path.splitext(name)[0] if re.search(r"\.[A-Za-z0-9]{2,5}$", name) else name
    base = re.sub(r"[_]+", " ", base)
    return re.sub(r"\s+", " ", base).strip() or name


# --------------------------------------------------------------------------
# Driver
# --------------------------------------------------------------------------


def sources():
    """Yield (source_path, base_meta) for everything that should end up in corpus/."""
    for entry in load_manifest().values():
        path = os.path.join(RAW_DIR, entry["path"])
        if not os.path.exists(path):
            continue
        name = entry.get("title") or os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()
        meta = {
            "title": clean_title(name) if entry["kind"] == "file" else name,
            "course": entry.get("course"), "course_name": entry.get("course_name"),
            "term": entry.get("term"), "module": entry.get("module"),
            "date": entry.get("date"), "source": "canvas", "url": entry.get("url"),
            "original": "raw/" + entry["path"],
        }
        fixed = KIND_TO_TYPE.get(entry["kind"])
        yield path, meta, (fixed or (lambda info, e=entry, x=ext: infer_type(
            e.get("title", ""), e.get("folder", ""), e.get("module", ""), x, info)))

    if not os.path.isdir(INBOX_DIR):
        return
    for dirpath, _, files in os.walk(INBOX_DIR):
        rel = os.path.relpath(dirpath, INBOX_DIR).replace(os.sep, "/")
        top = rel.split("/")[0]
        for fname in sorted(files):
            if fname.startswith(".") or fname.lower() == "readme.md":
                continue
            path = os.path.join(dirpath, fname)
            ext = os.path.splitext(fname)[1].lower()
            m = DATED.match(os.path.splitext(fname)[0])
            meta = {"title": clean_title(m.group(2) if m else fname),
                    "date": m.group(1) if m else "", "source": "upload",
                    "original": "inbox/" + os.path.relpath(path, INBOX_DIR).replace(os.sep, "/")}
            if top == "canvas" and "/" in rel:
                meta["course"] = rel.split("/")[1]
                meta["source"] = "canvas (manual)"
                kind = (lambda info, f=fname, x=ext: infer_type(f, "", "", x, info))
            else:
                kind = INBOX_TYPES.get(top, "note")
                if top not in INBOX_TYPES:
                    meta["module"] = rel if rel != "." else ""
            yield path, meta, kind


def output_path(meta):
    course = slugify(meta.get("course") or "")
    group = course if meta.get("course") else "_" + PLURAL[meta["type"]]
    sub = PLURAL[meta["type"]] if meta.get("course") else ""
    stem = slugify(((meta.get("date") or "") + " " + meta["title"]).strip()
                   if meta["type"] in ("announcement", "transcript", "newsletter", "note")
                   else meta["title"])
    return os.path.join(CORPUS_DIR, group, sub, stem + ".md")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--all", action="store_true", help="reconvert everything")
    args = ap.parse_args()

    state = {}
    if os.path.exists(STATE) and not args.all:
        with open(STATE, encoding="utf-8") as f:
            state = json.load(f)
    new_state, failures, claimed = {}, [], {}
    done = unchanged = 0

    for path, meta, kind in sources():
        key = os.path.relpath(path, ROOT).replace(os.sep, "/")
        digest = file_sha1(path)
        prev = state.get(key)
        if prev and prev["sha1"] == digest and os.path.exists(os.path.join(ROOT, prev["out"])):
            new_state[key] = prev
            claimed[prev["out"]] = key
            unchanged += 1
            continue
        ext = os.path.splitext(path)[1].lower()
        conv = CONVERTERS.get(ext)
        try:
            if not conv:
                raise Unsupported(f"{ext or 'no extension'} files aren't converted")
            body, info = conv(path)
        except Unsupported as e:
            failures.append({"file": key, "reason": str(e)})
            continue
        except Exception as e:  # noqa: BLE001 -- one bad file shouldn't stop the run
            failures.append({"file": key, "reason": f"{type(e).__name__}: {e}"})
            continue

        extra = info.get("meta") or {}
        for field in ("title", "course", "course_name", "term", "module", "date", "url"):
            if extra.get(field) and not meta.get(field):
                meta[field] = extra[field]
        if ext == ".eml" and info.get("title"):          # an email's subject beats its filename
            meta["title"] = info["title"].strip()
            if not meta.get("date") and info.get("date_header"):
                try:
                    meta["date"] = email.utils.parsedate_to_datetime(info["date_header"]).date().isoformat()
                except (TypeError, ValueError):
                    pass
        meta["type"] = kind(info) if callable(kind) else kind
        meta["locator_kind"] = info.get("locator_kind", "section")
        if not body.strip():
            failures.append({"file": key, "reason": "no text after conversion"})
            continue

        out = output_path(meta)
        rel_out = os.path.relpath(out, ROOT).replace(os.sep, "/")
        if rel_out in claimed and claimed[rel_out] != key:     # same title, different file
            out = out[:-3] + "-" + digest[:6] + ".md"
            rel_out = os.path.relpath(out, ROOT).replace(os.sep, "/")
        claimed[rel_out] = key
        if prev and prev["out"] != rel_out and os.path.exists(os.path.join(ROOT, prev["out"])):
            os.remove(os.path.join(ROOT, prev["out"]))
        write_doc(out, meta, body)
        new_state[key] = {"sha1": digest, "out": rel_out}
        done += 1

    # Sources that disappeared take their Markdown with them.
    for key, prev in state.items():
        if key not in new_state and prev["out"] not in claimed:
            target = os.path.join(ROOT, prev["out"])
            if os.path.exists(target):
                os.remove(target)

    os.makedirs(RAW_DIR, exist_ok=True)
    with open(STATE, "w", encoding="utf-8") as f:
        json.dump(new_state, f, indent=1, sort_keys=True)
    with open(CONVERT_FAILED, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["file", "reason"])
        w.writeheader()
        w.writerows(failures)

    print(f"Converted {done}, unchanged {unchanged}, could not convert {len(failures)}.")
    if failures:
        print(f"See {os.path.relpath(CONVERT_FAILED, ROOT)}. Common fixes: export scanned PDFs "
              "with OCR, or re-save old .doc/.ppt files.")
    print("Next: git add corpus && git commit && git push  (the app re-indexes on push)")


if __name__ == "__main__":
    main()
