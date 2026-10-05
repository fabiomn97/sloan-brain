"""Synthetic course materials for tests and local previews.

Written for this repo; not real Sloan materials. Builds PPTX, PDF, DOCX and
XLSX files so every converter path gets exercised.
"""

import os

PRICING_SLIDES = [
    ("Pricing and elasticity", ["15.010 Session 4", "What should we charge, and why?"]),
    ("Price elasticity of demand", [
        "Elasticity = % change in quantity / % change in price",
        "Elastic (|e| > 1): a price cut raises revenue",
        "Inelastic (|e| < 1): a price increase raises revenue",
        "Estimate it from experiments, not from intuition"]),
    ("The markup rule", [
        "Profit-maximizing markup: (P - MC) / P = 1 / |e|",
        "The less elastic the demand, the higher the margin you can sustain",
        "Marginal cost, not average cost, sets the floor"]),
    ("Price discrimination", [
        "First degree: charge each customer their willingness to pay",
        "Second degree: menus, versions, quantity discounts; customers self-select",
        "Third degree: different prices for observable groups (students, regions)",
        "Requires market power, segmentation, and no resale"]),
    ("Value-based pricing in practice", [
        "Start from the customer's next best alternative",
        "Quantify the differentiation value you add on top of it",
        "Price to share that value, not to cover cost-plus"]),
    ("Common mistakes", [
        "Cost-plus pricing ignores willingness to pay",
        "Matching competitor cuts starts price wars in commodity markets",
        "Discounting to hit volume targets trains customers to wait"]),
]

FEEDBACK_SLIDES = [
    ("Giving and receiving feedback", ["15.311 Organizational Processes, Session 7"]),
    ("Why feedback fails", [
        "We give conclusions, not observations",
        "The receiver hears a threat to identity and stops listening",
        "Feedback arrives too late to act on"]),
    ("The SBI model", [
        "Situation: when and where it happened",
        "Behavior: what the person did, observable and specific",
        "Impact: what it caused for you, the team, or the client",
        "Then ask, do not tell: 'How did you see it?'"]),
    ("Receiving feedback well", [
        "Separate the what from the who and the how it was delivered",
        "Ask for one specific example before reacting",
        "Thank, reflect, decide: you own what to change"]),
    ("Team norms", [
        "Agree in the team charter how and when you will give feedback",
        "Schedule a mid-semester feedback session before problems appear"]),
]

PRICE_DISCRIMINATION_NOTE = [
    "Note on Price Discrimination\n\nPrepared for 15.010 Economic Analysis for Business Decisions.\n\n"
    "Firms with market power rarely charge everyone the same price. Price discrimination lets a firm "
    "capture more of the consumer surplus that a single price leaves on the table. Airlines, software "
    "companies and movie theaters all use it, and the managerial question is never whether it is "
    "possible but which form fits the market.",
    "Three conditions must hold. First, the firm needs market power: in a perfectly competitive market "
    "any premium is competed away. Second, the firm must be able to sort customers by willingness to pay, "
    "either directly through observable traits or indirectly by letting them self-select. Third, resale "
    "must be difficult, otherwise low-price buyers become arbitrageurs.\n\n"
    "Versioning is the most common form of second-degree discrimination. A software firm sells a basic "
    "and a professional edition; the basic edition is deliberately limited so that high-value customers "
    "prefer to pay for the professional one. The design question is how much to degrade the low version.",
    "Managers should test segmentation ideas with small experiments, watch for fairness backlash when "
    "prices are visible to customers, and remember that the relevant comparison is against the next best "
    "alternative each segment has, not against cost.",
]

LITTLES_LAW = [
    ("Heading 1", "Little's Law and queueing"),
    (None, "15.761 Operations Management, class notes for Session 3."),
    ("Heading 2", "Little's Law"),
    (None, "Inventory = Throughput x Flow time (I = R x T). It holds for any stable process, "
           "whatever the arrival pattern or the order in which work is served."),
    (None, "Use it to find the missing number: if a clinic sees 30 patients an hour and the average "
           "patient spends 40 minutes inside, there are on average 20 patients in the building."),
    ("Heading 2", "Why queues explode near full utilization"),
    (None, "Waiting time grows non-linearly as utilization approaches 100%. Variability in arrivals "
           "and in service times makes it worse. Cutting variability is often cheaper than adding capacity."),
    ("List Bullet", "Pool capacity: one shared line beats separate lines per server."),
    ("List Bullet", "Protect the bottleneck: an hour lost there is an hour lost for the whole system."),
]

TRANSCRIPT = """Coffee chat with a product manager at a fintech company

Notes, 2 October 2026.

They moved from consulting into product three years ago. The advice that stuck:
- Recruiters screen for evidence that you shipped something, not for the title. Lead with the
  product decision you made and the metric that moved.
- For APM-style interviews, practice product sense out loud with a timer; the structure matters
  less than making a clear call and naming the tradeoff.
- Reach out to people two or three years ahead of you, not to senior leaders. They remember the process.
- Follow up within 24 hours with one specific thing you will do with their advice.
"""

NEWSLETTER = """From: Sloan Student Life <sloan-life@example.edu>
Subject: Sloan Weekly: negotiation workshop, career treks, and the Fall formal
Date: Mon, 28 Sep 2026 08:00:00 -0400
Content-Type: text/html; charset=utf-8

<html><body><h2>This week at Sloan</h2>
<p><strong>Negotiation workshop.</strong> Learn to prepare your BATNA and anchor first in salary
conversations. Thursday, 5:30pm, E62-233.</p>
<ul><li>Career treks to San Francisco open for applications until Oct 9.</li>
<li>Fall formal tickets go on sale Friday.</li></ul>
<p>Read more on the <a href="https://example.edu/weekly">student life site</a>.</p>
</body></html>
"""


def make_pptx(path, slides, notes=True):
    from pptx import Presentation
    prs = Presentation()
    for i, (title, bullets) in enumerate(slides):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if notes and i == 2:
            s.notes_slide.notes_text_frame.text = "Work the example: at e = -2 the markup is 50% of price."
    prs.save(path)


def make_pdf(path, pages, landscape=False, footer=None):
    import pymupdf
    doc = pymupdf.open()
    for text in pages:
        w, h = (792, 612) if landscape else (612, 792)
        page = doc.new_page(width=w, height=h)
        page.insert_textbox(pymupdf.Rect(60, 60, w - 60, h - 80), text, fontsize=13 if landscape else 11)
        if footer:
            page.insert_text((60, h - 40), footer, fontsize=8)
    doc.save(path)


def make_slide_pdf(path, slides):
    make_pdf(path, [t + "\n\n" + "\n".join("- " + b for b in bs) for t, bs in slides],
             landscape=True, footer="MIT Sloan School of Management")


def make_scanned_pdf(path):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 500, 700), color=(0.2, 0.2, 0.2), fill=(0.9, 0.9, 0.9))
    doc.save(path)


def make_docx(path, paras):
    import docx
    d = docx.Document()
    for style, text in paras:
        d.add_paragraph(text, style=style) if style else d.add_paragraph(text)
    d.save(path)


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Newsvendor"
    ws.append(["Input", "Value"])
    ws.append(["Price", 10])
    ws.append(["Unit cost", 4])
    ws.append(["Salvage value", 1])
    ws.append(["Critical ratio = (p - c) / (p - s)", "=(B2-B3)/(B2-B4)"])
    junk = wb.create_sheet("rsklibSimData")          # Analytic Solver scratch data
    junk.append(["_x0001__x0002_" * 2000])
    junk.sheet_state = "hidden"
    wb.save(path)


def build(out_dir):
    """Write every binary fixture into out_dir and return their paths by key."""
    os.makedirs(out_dir, exist_ok=True)
    paths = {
        "pricing_pptx": os.path.join(out_dir, "Session 4 - Pricing and Elasticity.pptx"),
        "pd_note_pdf": os.path.join(out_dir, "Note on Price Discrimination.pdf"),
        "feedback_pdf": os.path.join(out_dir, "Session 7 Feedback slides.pdf"),
        "littles_docx": os.path.join(out_dir, "Littles Law class notes.docx"),
        "newsvendor_xlsx": os.path.join(out_dir, "Newsvendor model.xlsx"),
        "scanned_pdf": os.path.join(out_dir, "Scanned handout.pdf"),
    }
    make_pptx(paths["pricing_pptx"], PRICING_SLIDES)
    make_pdf(paths["pd_note_pdf"], PRICE_DISCRIMINATION_NOTE, footer="15.010 Fall 2026 | Not for distribution")
    make_slide_pdf(paths["feedback_pdf"], FEEDBACK_SLIDES)
    make_docx(paths["littles_docx"], LITTLES_LAW)
    make_xlsx(paths["newsvendor_xlsx"])
    make_scanned_pdf(paths["scanned_pdf"])
    return paths
